#!/usr/bin/env python3
"""Build trust-phar-01-medicine-import-sample.xlsx from CSV + field guide sheet."""
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("Install openpyxl: pip install openpyxl")

ROOT = Path(__file__).resolve().parent
CSV_PATH = ROOT / "trust-phar-01-medicine-import-sample.csv"
OUT_PATH = ROOT / "trust-phar-01-medicine-import-sample.xlsx"
UI_PATH = ROOT.parent.parent / "shop-management-ui" / "src" / "assets" / "templates" / "trust-phar-01-medicine-import-sample.xlsx"

FIELD_GUIDE = [
    ("Column", "Import", "Required", "Format / rules", "Example"),
    ("Product Name", "Yes", "Yes", "Text", "Paracetamol 500mg Tablet"),
    ("Salt Composition", "Yes", "No", "Text → medicine composition", "Paracetamol 500 mg"),
    ("Product Code / SKU", "Yes", "Recommended", "Unique per outlet; short codes OK", "PARA-500"),
    ("Barcode", "Yes", "No", "Unique per outlet", "8901001002001"),
    ("Category", "Yes", "Yes", "Therapeutic / product class", "Analgesic & Antipyretic"),
    ("Brand", "Yes", "No", "Text", "Generic Pharma"),
    ("HSN Code", "Yes", "No", "4-8 digits", "30049099"),
    ("GST %", "Yes", "Yes", "0, 3, 5, 12, 18, or 28", "12"),
    ("Purchase Price", "Validated", "No", "Number >= 0", "18.00"),
    ("Selling Price", "Yes", "Yes", "Number > 0", "24.00"),
    ("MRP", "Yes", "No", "Number >= 0", "30.00"),
    ("Current Stock", "Yes", "Yes", "Integer >= 0", "500"),
    ("Unit Type", "Yes", "No", "strip, tablet, bottle, injection, etc.", "strip"),
    ("Batch Number", "Template", "No", "Use inventory batches after import", "B-PARA-001"),
    ("Expiry Date", "Template", "No", "YYYY-MM-DD or DD-MM-YYYY", "2027-08-31"),
    ("Manufacturer", "Yes", "No", "Text", "Micro Labs Ltd"),
    ("Minimum Stock Alert", "Yes", "No", "Integer >= 1", "50"),
    ("Product Status", "Yes", "No", "ACTIVE or INACTIVE", "ACTIVE"),
    ("Schedule Type", "Yes", "No", "H, H1, X, OTC, etc.", "H"),
    ("Generic Name", "No", "No", "Reference only", "Paracetamol"),
    ("Therapeutic Class", "No", "No", "Use Category for import", "Analgesic"),
    ("Supplier / Vendor", "No", "No", "Reference only", "MedSupply Distributors"),
    ("Manufacturing Date", "No", "No", "Reference only", "2025-11-01"),
    ("Pack Size", "No", "No", "Reference only", "10x10 strips"),
    ("Discount %", "No", "No", "Reference only", "0"),
    ("Rack / Shelf Location", "No", "No", "Reference only", "A-01-R1"),
    ("Prescription Required", "No", "No", "Reference only", "Yes/No"),
    ("Dosage Form", "No", "No", "Use Unit Type hint", "Tablet"),
    ("Strength", "No", "No", "Include in Product Name", "500 mg"),
    ("Storage Condition", "No", "No", "Reference only", "Below 30 C"),
    ("Medicine Image URL", "No", "No", "Not supported", ""),
    ("Notes / Remarks", "No", "No", "Reference only", "Fast-moving OTC"),
]


def load_csv_rows():
    import csv
    with CSV_PATH.open(encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        return list(reader)


def style_header_row(ws, row_idx=1):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[row_idx]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def auto_width(ws, max_width=48):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[letter].width = min(max(length + 2, 10), max_width)


def main():
    rows = load_csv_rows()
    wb = Workbook()
    ws_data = wb.active
    ws_data.title = "Medicines"
    for r in rows:
        ws_data.append(r)
    style_header_row(ws_data)
    ws_data.freeze_panes = "A2"
    auto_width(ws_data)

    ws_guide = wb.create_sheet("Field guide")
    for r in FIELD_GUIDE:
        ws_guide.append(r)
    style_header_row(ws_guide)
    for row in ws_guide.iter_rows(min_row=2, max_row=ws_guide.max_row):
        imp = row[1].value
        if imp == "Yes":
            row[1].fill = PatternFill("solid", fgColor="C6EFCE")
        elif imp == "Validated":
            row[1].fill = PatternFill("solid", fgColor="FFEB9C")
        else:
            row[1].fill = PatternFill("solid", fgColor="F2F2F2")
        req = row[2].value
        if req == "Yes":
            row[2].font = Font(bold=True, color="C00000")
    auto_width(ws_guide, 60)

    wb.save(OUT_PATH)
    UI_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(UI_PATH)
    print(f"Wrote {OUT_PATH}")
    print(f"Wrote {UI_PATH}")


if __name__ == "__main__":
    main()
